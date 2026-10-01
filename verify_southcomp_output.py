from sales.southcomp_engine import generate_item_creation_excel
from openpyxl import load_workbook
import io

rows = [('3400023336849.1-210-BPCK', 'Dell Pro 16 Plus 16" Cu7-265U 12C 5.3GHz 16GB 512GB SSD Wi-Fi 6E AX211 Windows 11 Pro')]
b = generate_item_creation_excel(rows)
wb = load_workbook(io.BytesIO(b), data_only=True)
ws = wb.active
print('rows', ws.max_row, 'cols', ws.max_column)
for i in range(1, 4):
    vals = [ws.cell(i, c).value for c in range(1, ws.max_column + 1)]
    print('ROW', i, vals[:10])
print('A2', ws['A2'].value)
print('B2', ws['B2'].value)
print('A3', ws['A3'].value)
