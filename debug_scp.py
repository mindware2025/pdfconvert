from openpyxl import load_workbook

wb = load_workbook(r'C:\Users\Z.Mama\pdfconvert\SCP-IT20260922-SAMPLE DELL LAPTOP .xlsx', data_only=True)
ws = wb.active
print('max_row', ws.max_row, 'max_col', ws.max_column)
for i in range(1, min(ws.max_row + 1, 4)):
    vals = [ws.cell(i, c).value for c in range(1, ws.max_column + 1)]
    print('ROW', i, len(vals))
    print(vals)
